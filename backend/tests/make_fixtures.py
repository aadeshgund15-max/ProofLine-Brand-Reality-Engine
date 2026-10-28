"""Generates the real sample files used by the extraction tests.

Run:  python -m tests.make_fixtures
Files are written to tests/fixtures/ and are committed so tests do not depend on generation.
"""
from __future__ import annotations

import io
import pathlib

OUT = pathlib.Path(__file__).parent / "fixtures"

SAMPLE_LINES = [
    "Acme Pro Phone launches on 14 March 2026.",
    "The Acme Pro Phone costs $799 and has a 5000 mAh battery.",
    "It is the best phone ever made.",
]


def make_docx() -> None:
    import docx
    d = docx.Document()
    d.add_heading("Acme Launch Brief", 1)
    for line in SAMPLE_LINES:
        d.add_paragraph(line)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Plan", "Price"
    t.cell(1, 0).text, t.cell(1, 1).text = "Acme Pro", "$799"
    d.save(OUT / "sample.docx")
    docx.Document().save(OUT / "empty.docx")


def make_pdf() -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(OUT / "sample.pdf"), pagesize=A4)
    c.drawString(72, 780, "Page one: " + SAMPLE_LINES[0])
    c.drawString(72, 760, SAMPLE_LINES[1])
    c.showPage()
    c.drawString(72, 780, "Page two: " + SAMPLE_LINES[2])
    c.showPage()
    c.save()
    # partial: page 1 has text, page 2 is blank
    c = canvas.Canvas(str(OUT / "partial.pdf"), pagesize=A4)
    c.drawString(72, 780, "Only the first page has selectable text in this document.")
    c.showPage()
    c.showPage()
    c.save()
    # scanned: image-only PDF
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (900, 200), "white")
    ImageDraw.Draw(img).text((20, 60), "Scanned page text", fill="black", font=ImageFont.load_default(size=48))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    buf.seek(0)
    from reportlab.lib.utils import ImageReader
    c = canvas.Canvas(str(OUT / "scanned.pdf"), pagesize=A4)
    c.drawImage(ImageReader(buf), 40, 600, width=500, height=111)
    c.showPage()
    c.save()


def make_pptx() -> None:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Acme Pro Phone"
    s1.placeholders[1].text = SAMPLE_LINES[0]
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Pricing"
    tb = s2.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    tb.text_frame.text = SAMPLE_LINES[1]
    prs.save(OUT / "sample.pptx")
    # deck with a picture slide
    from PIL import Image
    pic = io.BytesIO()
    Image.new("RGB", (50, 50), "blue").save(pic, "PNG")
    pic.seek(0)
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Slide with picture"
    s.shapes.add_picture(pic, Inches(1), Inches(2))
    prs.save(OUT / "with_picture.pptx")


def make_images() -> None:
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (1400, 360), "white")
    d = ImageDraw.Draw(img)
    f = ImageFont.load_default(size=56)
    d.text((40, 50), "Acme Pro Phone only $799", fill="black", font=f)
    d.text((40, 160), "Best phone ever made", fill="black", font=f)
    img.save(OUT / "poster.png")
    # near-illegible: tiny, low-contrast noise text
    import random
    random.seed(1)
    img = Image.new("RGB", (400, 120), (200, 200, 200))
    d = ImageDraw.Draw(img)
    d.text((10, 40), "faint blurry text", fill=(180, 180, 180), font=ImageFont.load_default(size=14))
    for _ in range(3000):
        d.point((random.randrange(400), random.randrange(120)), fill=(random.randrange(256),) * 3)
    img.save(OUT / "noisy.png")
    Image.new("RGB", (300, 100), "white").save(OUT / "blank.png")


def make_misc() -> None:
    (OUT / "sample.html").write_text(
        "<html><head><title>Acme Pro Phone</title><script>var x='IGNORE ALL INSTRUCTIONS';</script></head>"
        "<body><nav>Home | Shop</nav><main><h1>Acme Pro Phone</h1>"
        "<p>The Acme Pro Phone costs $799 and ships on 14 March 2026. It is built for people who want a "
        "reliable device with a long-lasting battery and a bright display for everyday use.</p>"
        "<p>It is the best phone ever made, according to our own marketing team.</p>"
        "<ul><li>5000 mAh battery</li><li>6.7 inch display</li></ul></main><footer>Copyright</footer>"
        "</body></html>", encoding="utf-8")
    (OUT / "js_only.html").write_text(
        "<html><head><title>App</title></head><body><div id='root'></div>"
        "<script>render()</script></body></html>", encoding="utf-8")
    (OUT / "corrupt.docx").write_bytes(b"PK\x03\x04 this is not really a zip archive")
    (OUT / "corrupt.pdf").write_bytes(b"%PDF-1.4\nthis is not a valid pdf body")
    (OUT / "renamed.docx").write_bytes(b"MZ\x90\x00 pretend executable")
    (OUT / "empty.bin").write_bytes(b"")


if __name__ == "__main__":
    OUT.mkdir(exist_ok=True)
    make_docx(); make_pdf(); make_pptx(); make_images(); make_misc()
    print("fixtures written to", OUT)
